from datetime import date
from decimal import Decimal, InvalidOperation
from pathlib import Path

from django.conf import settings
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from apps.companies.models import Company
from apps.employees.models import Employee
from apps.payroll.models import EmployeeSalary


SOURCE_ROWS = (
    ("ST", "D RAVI", "100000"),
    ("ST", "C MURUGAIAN", "40000"),
    ("ST", "DHANAPAL", "50000"),
    ("ST", "PRABAKAR", "18000"),
    ("ST", "PANDIARAJAN", "20000"),
    ("ST", "V N IYYAPPAN", "30000"),
    ("ST", "KARTHICK", "25000"),
    ("ST", "SUNDARAM", "25000"),
    ("ST", "VASU", "24000"),
    ("ST", "SEKAR", "15000"),
    ("ST", "ARUNKUMAR", "15000"),
    ("ST", "THULAL BINDHANI", "20000"),
    ("ST", "JOJU", "17000"),
    ("ST", "RAJA BINDHANI", "17000"),
    ("ST", "MAHENDRAN", "18000"),
    ("ST", "MANGAL", "17000"),
    ("ST", "CHINNARAJ", "50000"),
    ("ST", "SATHYA SELAN", "25000"),
    ("ST", "REGAN", "16000"),
    ("ST", "PAPPU", "18000"),
    ("ST", "AJAY", "18000"),
    ("ST", "JANARTHANAN BANDIA", "15000"),
    ("ST", "PRITHILAL", "15000"),
    ("ST", "SAM BIHARI", "15000"),
    ("ST", "ANIL", "18000"),
    ("ST", "RAJAN LOHAR", "15000"),
    ("ST", "TIRTH KARSHAL", "15000"),
    ("ST", "RANJANI", "12000"),
    ("SUT", "D RAVI", "100000"),
    ("SUT", "S SUNDARAMOORTHY", "50000"),
    ("SUT", "RAMACHANDRA", "22000"),
    ("SUT", "CHAMPAI SOY", "15500"),
    ("SUT", "JAYAPAL", "17000"),
    ("SUT", "MAN SINGH", "15500"),
    ("SUT", "RAGHUL", "15000"),
    ("SUT", "PARAMESWAR", "15000"),
    ("SUT", "LOL MOHAN SOY", "14500"),
    ("SUT", "NIKIL", "21000"),
    ("SUT", "CHOTTU", "15000"),
    ("SUT", "DEEPAK", "21000"),
    ("SUT", "RAM JANMA", "21000"),
    ("SUT", "SHANKAR", "15000"),
    ("SUT", "RADHE SHYAM", "15000"),
    ("SUT", "RAMO", "15000"),
    ("SUT", "PARVIN", "15000"),
    ("SUT", "SURAP MARK", "15000"),
    ("SUT", "GURU CHARAN", "15000"),
    ("SUT", "SUKULAL", "15000"),
    ("SUT", "BHURBADRA", "15500"),
    ("SUT", "CHANDAN", "14500"),
    ("SUT", "MADHU", "14500"),
    ("SUT", "TINTU", "14500"),
    ("SUT", "EGAVALLI", "13000"),
    ("SBT", "RAVI", "100000"),
    ("SBT", "SOUNDARA PANDI", "60000"),
    ("SBT", "BHARATHIDASAN", "50000"),
    ("SBT", "AJITH", "27000"),
    ("SBT", "RANJITH Q", "23000"),
    ("SBT", "SECURITY", "15000"),
    ("SBT", "KESHAP", "20000"),
    ("SBT", "SUSHIL", "18000"),
    ("SBT", "TRIPATI", "18000"),
    ("SBT", "CHANDAN", "18000"),
    ("SBT", "BHARATH", "15000"),
    ("SBT", "MAHA", "22000"),
    ("SBT", "KISKU", "14000"),
    ("SBT", "DURUGHA", "14000"),
    ("SBT", "KABIN", "30000"),
    ("SBT", "UPENDHAR", "27500"),
    ("SBT", "LAKSHMI DHAR", "23500"),
    ("SBT", "JANAKIRAMAN", "14500"),
    ("SBT", "SOMNATH SOY", "16000"),
    ("SBT", "SUNARAM", "15500"),
    ("SBT", "DIBARA PURTHY", "15000"),
    ("SBT", "GANESH", "15500"),
    ("SBT", "TRILOCHAN", "16000"),
    ("SBT", "SAMU", "16000"),
    ("SBT", "MANSING", "15000"),
    ("SBT", "RANJITH", "16500"),
    ("SBT", "MARKANDA", "14500"),
    ("SBT", "BEEMSON PALIYA", "16000"),
    ("SBT", "KAMAL", "14500"),
    ("SBT", "VAIRAB", "14500"),
    ("SBT", "RABI", "14500"),
    ("SBT", "LAL MUNDA", "16500"),
    ("SBT", "ANIRUTH", "23000"),
    ("SBT", "LADURAM", "15500"),
    ("SBT", "GANBINDO PANSING", "14500"),
    ("SBT", "JADOP SOY", "14500"),
    ("SBT", "DEBO", "14500"),
    ("SBT", "RASANANDA", "14500"),
    ("SBT", "RAMAKRISHNA", "14500"),
    ("SBT", "SUGUN SOY", "15500"),
    ("SBT", "DULU SOY", "15000"),
    ("SBT", "KRISHNA", "15500"),
    ("SBT", "JITEN", "14500"),
    ("SBT", "HAJAN", "14500"),
    ("SBT", "RAJINESH", "14500"),
    ("SBT", "ANUBHAB", "14500"),
    ("SBT", "JAGA MOHAN", "16000"),
    ("SBT", "FAKHIRMAN", "15500"),
    ("SBT", "KANDE HO", "14500"),
    ("SBT", "HINDU SOY", "14500"),
    ("SBT", "PRAVATH", "14500"),
    ("SBT", "KANIEI", "15500"),
    ("SBT", "DHAASEKAR", "14500"),
    ("SBT", "RANJITH II", "14500"),
    ("SBT", "UMASH", "14500"),
    ("SBT", "SRIRAM", "15000"),
    ("SBT", "SHARKAR", "14500"),
    ("SBT", "CHANDAN PINGUA", "15000"),
    ("SBT", "SANKAR", "14500"),
    ("SBT", "LOURTHU SAMY", "22000"),
    ("SBT", "RATTAN", "20000"),
    ("SBT", "DEBENTHRA", "18000"),
    ("SBT", "AKASH", "25000"),
    ("SBT", "SAANTI MAAJI", "18000"),
    ("SBT", "RENUKA", "13000"),
    ("SBT", "CHRCHAMMAL", "12000"),
    ("SBT", "POONKODI", "12000"),
    ("SBT", "BAIRAJUL ANSARI", "14500"),
)


def normalize_name(value):
    return " ".join(value.split()).casefold()


class Command(BaseCommand):
    help = "Safely import current monthly salaries for existing employees"

    def add_arguments(self, parser):
        parser.add_argument(
            "--workbook",
            default=str(Path(settings.BASE_DIR).parent / "JULY 2026_ST.xlsx"),
            help="Salary workbook path",
        )
        parser.add_argument(
            "--dry-run",
            action="store_true",
            help="Report changes without modifying the database",
        )
        parser.add_argument(
            "--effective-from",
            type=date.fromisoformat,
            default=date.today,
            help="Effective date for newly created salary structures (YYYY-MM-DD)",
        )

    def handle(self, *args, **options):
        dry_run = options["dry_run"]
        effective_from = options["effective_from"]
        source_rows = self._load_source_rows(options["workbook"])
        companies = self._get_companies()
        results = self._classify(companies, source_rows)

        counts = {"matched": 0, "unchanged": 0, "updated": 0, "unmatched": 0, "ambiguous": 0}
        for result in results:
            status = result["status"]
            counts[status] += 1
            if status in {"unmatched", "ambiguous"}:
                candidates = ", ".join(
                    employee.employee_code
                    for employee in result["candidates"]
                )
                suffix = f" | Candidates: {candidates}" if candidates else ""
                self.stdout.write(
                    f"{status.upper()}: {result['company']} | {result['name']}{suffix}"
                )
                continue
            current = result["current"]
            new_salary = result["salary"]
            action = "UNCHANGED" if current == new_salary else "UPDATED"
            self.stdout.write(
                f"{action}: {result['company']} | {result['employee'].employee_code} | "
                f"{result['name']} | Old Salary: {current or 'NONE'} | New Salary: {new_salary}"
            )
            if action == "UNCHANGED":
                counts["unchanged"] += 1
            else:
                counts["updated"] += 1

        self.stdout.write("")
        self.stdout.write(
            "MATCHED: {matched}\nUNCHANGED: {unchanged}\nUPDATED: {updated}\n"
            "UNMATCHED: {unmatched}\nAMBIGUOUS: {ambiguous}".format(**counts)
        )

        if dry_run:
            self.stdout.write(self.style.WARNING("DRY RUN: no database changes made"))
            return

        if counts["unmatched"] or counts["ambiguous"]:
            self.stdout.write(
                self.style.WARNING(
                    f"{counts['unmatched'] + counts['ambiguous']} rows were skipped "
                    "(unmatched/ambiguous) and will NOT be imported."
                )
            )

        with transaction.atomic():
            created = 0
            for result in results:
                if result["status"] != "matched":
                    continue
                if result["current"] == result["salary"]:
                    continue
                salary, was_created = EmployeeSalary.objects.get_or_create(
                    employee=result["employee"],
                    is_active=True,
                    defaults={
                        "basic_salary": result["salary"],
                        "effective_from": effective_from,
                    },
                )
                if was_created:
                    created += 1
                if salary.basic_salary != result["salary"]:
                    salary.basic_salary = result["salary"]
                    salary.save(update_fields=["basic_salary"])

        self.stdout.write(
            self.style.SUCCESS(
                f"Salary import completed successfully ({created} created)"
            )
        )

    def _get_companies(self):
        companies = {}
        for code in {row[0] for row in SOURCE_ROWS}:
            try:
                companies[code] = Company.objects.get(code=code)
            except Company.DoesNotExist as exc:
                raise CommandError(f"Required company code does not exist: {code}") from exc
        return companies

    def _load_source_rows(self, workbook_path):
        path = Path(workbook_path)
        if not path.is_file():
            raise CommandError(f"Salary workbook does not exist: {path}")

        try:
            workbook = load_workbook(path, data_only=True, read_only=True)
        except Exception as exc:
            raise CommandError(f"Could not read salary workbook: {path}") from exc

        sheet_specs = {
            "SRINIVASA": ("ST", "name", "salary"),
            "SUDHARSHAN": ("SUT", "name", "salary"),
            "SBT": ("SBT", "name", "salary"),
        }
        rows = []
        for sheet_name, (company_code, name_header, salary_header) in sheet_specs.items():
            if sheet_name not in workbook.sheetnames:
                raise CommandError(f"Required workbook sheet does not exist: {sheet_name}")
            sheet = workbook[sheet_name]
            header_row = None
            name_column = None
            salary_column = None
            for row in sheet.iter_rows(values_only=True):
                headers = [
                    " ".join(str(value).split()).casefold()
                    if value is not None else ""
                    for value in row
                ]
                for index, header in enumerate(headers):
                    if header == name_header and name_column is None:
                        name_column = index
                    is_salary_header = (
                        header == salary_header
                        or (
                            salary_header == "salary"
                            and header.startswith("salary (in rs.)")
                        )
                    )
                    if is_salary_header and salary_column is None:
                        salary_column = index
                if name_column is not None and salary_column is not None:
                    header_row = row
                    break
            if header_row is None:
                raise CommandError(f"Could not find Name and Salary columns in {sheet_name}")

            for row in sheet.iter_rows(values_only=True):
                if len(row) <= max(name_column, salary_column):
                    continue
                name = row[name_column]
                salary = row[salary_column]
                if isinstance(name, str) and name.strip() and isinstance(salary, (int, float, Decimal)):
                    rows.append((company_code, name.strip(), str(salary)))
        return tuple(rows)

    def _match_employee(self, company, source_name, salary):
        """Resolve a source name to an employee within a company.

        Exact normalized match first; then order-insensitive token match when
        unambiguous; finally disambiguate known duplicate names by salary.
        Returns (employee or None, status).
        """
        name_key = normalize_name(source_name)
        employees = list(
            Employee.objects.filter(company=company).only(
                "id", "employee_code", "first_name", "last_name"
            )
        )

        def norm_name(emp):
            return normalize_name(f"{emp.first_name} {emp.last_name}")

        def disambiguate(candidates):
            """Return a single employee if resolvable, else (None, 'ambiguous')."""
            if len(candidates) == 1:
                return candidates[0], "matched"
            if len(candidates) > 1:
                salary_value = int(salary)
                by_salary = {
                    23000: "SBT001",
                    16500: "SBT024",
                }
                target_code = by_salary.get(salary_value)
                if target_code:
                    for emp in candidates:
                        if emp.employee_code == target_code:
                            return emp, "matched"
            return None, "ambiguous"

        exact = [emp for emp in employees if norm_name(emp) == name_key]
        if exact:
            return disambiguate(exact)

        source_tokens = {tok for tok in name_key.split() if len(tok) >= 2}
        if source_tokens:
            order_matches = [
                emp
                for emp in employees
                if source_tokens
                <= {tok for tok in norm_name(emp).split() if len(tok) >= 2}
            ]
            if order_matches:
                return disambiguate(order_matches)

        return None, "unmatched"

    def _classify(self, companies, source_rows):
        results = []
        for company_code, source_name, salary_text in source_rows:
            salary = self._parse_salary(salary_text, source_name)
            company = companies[company_code]
            employee, status = self._match_employee(
                company, source_name, salary
            )
            result = {
                "company": company_code,
                "name": source_name,
                "salary": salary,
                "employee": employee,
                "candidates": [employee] if employee else [],
                "current": None,
                "status": status,
            }
            if status == "matched":
                current = (
                    EmployeeSalary.objects.filter(
                        employee=employee, is_active=True
                    )
                    .order_by("-effective_from", "-id")
                    .first()
                )
                result["current"] = current.basic_salary if current else None
            results.append(result)
        return results

    def _parse_salary(self, value, source_name):
        try:
            salary = Decimal(value).quantize(Decimal("0.01"))
        except (InvalidOperation, TypeError, ValueError) as exc:
            raise CommandError(f"Invalid salary for {source_name}: {value}") from exc
        if salary < 0:
            raise CommandError(f"Salary cannot be negative for {source_name}: {salary}")
        return salary
