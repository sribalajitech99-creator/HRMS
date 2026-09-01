from datetime import datetime
from io import BytesIO

from django.http import HttpResponse

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import mm
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from reportlab.lib.styles import getSampleStyleSheet


def excel_filename(slug):
    today = datetime.now().strftime("%Y%m%d")
    return f"{slug}_{today}.xlsx"


def pdf_filename(slug):
    today = datetime.now().strftime("%Y%m%d")
    return f"{slug}_{today}.pdf"


def build_xlsx_response(
    columns,
    rows,
    title,
    filename="report.xlsx",
):
    workbook = Workbook()

    sheet = workbook.active

    sheet.title = title[:31] or "Report"

    sheet.append([title])

    sheet["A1"].font = Font(
        bold=True,
        size=14,
        color="FFFFFF",
    )

    sheet["A1"].fill = PatternFill(
        start_color="1B5E20",
        end_color="1B5E20",
        fill_type="solid",
    )

    sheet.append([])

    header = sheet.cell(
        row=3,
        column=1,
        value=title,
    )

    sheet.append(columns)

    header_row = sheet.max_row

    for cell in sheet[header_row]:
        cell.font = Font(
            bold=True,
            color="FFFFFF",
        )

        cell.fill = PatternFill(
            start_color="2E7D32",
            end_color="2E7D32",
            fill_type="solid",
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    for row in rows:
        clean = []

        for value in row:
            if value is None:
                clean.append("")
            elif isinstance(value, (datetime,)):
                clean.append(value.strftime("%Y-%m-%d %H:%M"))
            elif hasattr(value, "isoformat"):
                clean.append(value)
            elif isinstance(value, float):
                clean.append(round(value, 2))
            else:
                clean.append(value)

        sheet.append(clean)

    width = max(
        8,
        min(40, len(columns) * 2),
    )

    for index in range(1, len(columns) + 1):
        sheet.column_dimensions[
            get_column_letter(index)
        ].width = width

    for row in sheet.iter_rows(
        min_row=header_row + 1,
        max_col=len(columns),
    ):
        for cell in row:
            cell.alignment = Alignment(
                vertical="center",
            )

    sheet.freeze_panes = f"A{header_row + 1}"

    buffer = BytesIO()

    workbook.save(buffer)

    buffer.seek(0)

    response = HttpResponse(
        buffer,
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
    )

    response["Content-Disposition"] = (
        'attachment; filename="'
        f"{filename}\""
    )

    return response


def build_pdf_response(
    company_name,
    report_title,
    period,
    columns,
    rows,
    filename="report.pdf",
):
    buffer = BytesIO()

    styles = getSampleStyleSheet()

    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        rightMargin=10 * mm,
        leftMargin=10 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
        title=report_title,
    )

    story = []

    title_style = styles["Title"]
    title_style.textColor = colors.HexColor("#1B5E20")

    story.append(
        Paragraph(
            company_name or "HRMS",
            title_style,
        )
    )

    story.append(
        Paragraph(
            f"<b>{report_title}</b>",
            styles["Heading2"],
        )
    )

    story.append(
        Paragraph(
            f"Period: {period}",
            styles["Normal"],
        )
    )

    story.append(
        Paragraph(
            "Generated: "
            + datetime.now().strftime(
                "%d %b %Y, %H:%M"
            ),
            styles["Normal"],
        )
    )

    story.append(Spacer(1, 6 * mm))

    table_data = [columns]

    for row in rows:
        clean = []

        for value in row:
            if value is None:
                clean.append("")
            elif hasattr(value, "isoformat"):
                clean.append(
                    value.strftime(
                        "%Y-%m-%d"
                    )
                    if hasattr(value, "strftime")
                    else str(value)
                )
            else:
                clean.append(str(value))

        table_data.append(clean)

    table = Table(table_data, repeatRows=1)

    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2E7D32")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, 0), 8),
                ("FONTSIZE", (0, 1), (-1, -1), 7),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F8E9")]),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("ALIGN", (0, 0), (-1, -1), "LEFT"),
            ]
        )
    )

    story.append(table)

    doc.build(story)

    buffer.seek(0)

    response = HttpResponse(
        buffer,
        content_type="application/pdf",
    )

    response["Content-Disposition"] = (
        'attachment; filename="'
        f"{filename}\""
    )

    return response